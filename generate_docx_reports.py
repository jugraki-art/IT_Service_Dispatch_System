import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document_from_markdown(md_file_path, output_docx_path, doc_title, subtitle):
    doc = Document()
    
    # Page Setup - Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # #1e293b

    # Title Block
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(doc_title)
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(37, 99, 235) # Blue

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_p.add_run(subtitle)
    subtitle_run.font.size = Pt(13)
    subtitle_run.font.color.rgb = RGBColor(100, 116, 139) # Slate
    doc.add_paragraph() # Spacing

    with open(md_file_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []
    in_code_block = False
    code_lines = []

    for line in lines:
        stripped = line.rstrip()

        # Handle Code blocks
        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                # Output code box
                box_table = doc.add_table(rows=1, cols=1)
                box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = box_table.cell(0, 0)
                set_cell_background(cell, "F1F5F9")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                c_run = cp.add_run('\n'.join(code_lines))
                c_run.font.name = 'Consolas'
                c_run.font.size = Pt(9.5)
                c_run.font.color.rgb = RGBColor(15, 23, 42)
                doc.add_paragraph()
            continue

        if in_code_block:
            code_lines.append(stripped)
            continue

        # Handle Tables
        if '|' in stripped and ('|' in stripped[:3] or stripped.startswith('|')):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            # End of table
            in_table = False
            # Parse table_lines
            rows_data = []
            for t_line in table_lines:
                # Skip divider line like |---|---|
                parts = [p.strip() for p in t_line.split('|')]
                if len(parts) >= 2 and parts[0] == '' and parts[-1] == '':
                    parts = parts[1:-1]
                if all(set(p).issubset({'-', ':', ' '}) for p in parts if p):
                    continue
                rows_data.append(parts)

            if rows_data:
                num_cols = max(len(r) for r in rows_data)
                table = doc.add_table(rows=len(rows_data), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row in enumerate(rows_data):
                    for c_idx, val in enumerate(row):
                        if c_idx < num_cols:
                            c = table.cell(r_idx, c_idx)
                            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
                            p = c.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            run = p.add_run(val)
                            run.font.size = Pt(9.5)
                            if r_idx == 0:
                                set_cell_background(c, "E2E8F0")
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(15, 23, 42)
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(c, "F8FAFC")
                doc.add_paragraph()
            table_lines = []

        # Headings
        if stripped.startswith('# '):
            continue # Already handled in title
        elif stripped.startswith('## '):
            h2 = doc.add_heading(level=1)
            h2_run = h2.add_run(stripped[3:])
            h2_run.font.size = Pt(16)
            h2_run.font.bold = True
            h2_run.font.color.rgb = RGBColor(30, 64, 175) # #1e40af
            h2.paragraph_format.space_before = Pt(14)
            h2.paragraph_format.space_after = Pt(6)
        elif stripped.startswith('### '):
            h3 = doc.add_heading(level=2)
            h3_run = h3.add_run(stripped[4:])
            h3_run.font.size = Pt(13)
            h3_run.font.bold = True
            h3_run.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
            h3.paragraph_format.space_before = Pt(10)
            h3.paragraph_format.space_after = Pt(4)
        elif stripped.startswith('#### '):
            h4 = doc.add_heading(level=3)
            h4_run = h4.add_run(stripped[5:])
            h4_run.font.size = Pt(11)
            h4_run.font.bold = True
            h4_run.font.color.rgb = RGBColor(71, 85, 105)
            h4.paragraph_format.space_before = Pt(6)
            h4.paragraph_format.space_after = Pt(2)
        elif stripped.startswith('> '):
            # Blockquote / Callout
            quote_table = doc.add_table(rows=1, cols=1)
            quote_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            q_cell = quote_table.cell(0, 0)
            set_cell_background(q_cell, "EFF6FF")
            set_cell_margins(q_cell, top=100, bottom=100, left=160, right=160)
            qp = q_cell.paragraphs[0]
            qp.paragraph_format.space_before = Pt(2)
            qp.paragraph_format.space_after = Pt(2)
            q_run = qp.add_run(stripped[2:])
            q_run.font.size = Pt(10)
            q_run.font.italic = True
            q_run.font.color.rgb = RGBColor(30, 64, 175)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(stripped[2:])
        elif stripped.strip() == '---':
            # Separator
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run('—' * 45)
            run.font.color.rgb = RGBColor(203, 213, 225)
        elif stripped.strip() == '':
            pass
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.add_run(stripped)

    try:
        doc.save(output_docx_path)
        print(f"[OK] Generated document: {output_docx_path}")
    except PermissionError:
        print(f"[WARNING] Could not save {output_docx_path} (file is currently locked by another program like MS Word).")

def safe_copy(src, dst):
    try:
        shutil.copy(src, dst)
        print(f"[OK] Copied {os.path.basename(src)} -> {dst}")
    except Exception as e:
        print(f"[WARNING] Could not copy {os.path.basename(src)} to {dst}: {e}")

def main():
    base_dir = r"E:\xampp\htdocs\react-nest"
    docs_dir = r"C:\Users\gotten\Downloads\docs"

    doc1_md = os.path.join(base_dir, "SYSTEM_IMPLEMENTATION_AND_ADDITIONS_GUIDE.md")
    doc1_docx = os.path.join(base_dir, "SYSTEM_IMPLEMENTATION_AND_ADDITIONS_GUIDE.docx")
    try:
        create_document_from_markdown(
            doc1_md,
            doc1_docx,
            "IT Service Dispatch & Fair Odds Management System",
            "System Implementation & Technical Additions Guide"
        )
    except Exception as e:
        print(f"Error on doc1: {e}")

    doc2_md = os.path.join(base_dir, "IT_DISPATCH_WORKFLOW_AND_OPERATIONS_MANUAL.md")
    doc2_docx = os.path.join(base_dir, "IT_DISPATCH_WORKFLOW_AND_OPERATIONS_MANUAL.docx")
    try:
        create_document_from_markdown(
            doc2_md,
            doc2_docx,
            "IT Service Dispatch & Fair Odds Management System",
            "Operational Workflow & Lifecycle Architecture Manual"
        )
    except Exception as e:
        print(f"Error on doc2: {e}")

    doc3_md = os.path.join(base_dir, "PROJECT_ARCHITECTURE_ERD_AND_FILE_DIRECTORY.md")
    doc3_docx = os.path.join(base_dir, "PROJECT_ARCHITECTURE_ERD_AND_FILE_DIRECTORY.docx")
    try:
        create_document_from_markdown(
            doc3_md,
            doc3_docx,
            "IT Service Dispatch & Fair Odds Management System",
            "System Architecture, ERD & Complete File Directory Manual"
        )
    except Exception as e:
        print(f"Error on doc3: {e}")

    # Copy files to C:\Users\gotten\Downloads\docs
    safe_copy(doc1_md, os.path.join(docs_dir, "SYSTEM_IMPLEMENTATION_AND_ADDITIONS_GUIDE.md"))
    safe_copy(doc1_docx, os.path.join(docs_dir, "SYSTEM_IMPLEMENTATION_AND_ADDITIONS_GUIDE.docx"))
    safe_copy(doc2_md, os.path.join(docs_dir, "IT_DISPATCH_WORKFLOW_AND_OPERATIONS_MANUAL.md"))
    safe_copy(doc2_docx, os.path.join(docs_dir, "IT_DISPATCH_WORKFLOW_AND_OPERATIONS_MANUAL.docx"))
    safe_copy(doc3_md, os.path.join(docs_dir, "PROJECT_ARCHITECTURE_ERD_AND_FILE_DIRECTORY.md"))
    safe_copy(doc3_docx, os.path.join(docs_dir, "PROJECT_ARCHITECTURE_ERD_AND_FILE_DIRECTORY.docx"))

if __name__ == "__main__":
    main()
